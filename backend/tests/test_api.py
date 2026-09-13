"""End-to-end API workflow: project → PDF upload → review/correct → commit → pilot → enrichment → capture → grid → export."""

import io

import openpyxl
import respx

from app.jobs.worker import Worker
from tests.conftest import FIXTURES, MONTCO_PDF, gis_router, taxclaim_router


def run_jobs() -> int:
    return Worker(enable_scheduler=False).run_until_empty()


def test_full_pilot_workflow(client):
    assert client.get("/api/health").json()["status"] == "ok"
    meta = client.get("/api/meta").json()
    assert {c["county"] for c in meta["counties"]} == {"montco", "delco"}

    project = client.post("/api/projects", json={"name": "Montco 2026 upset sale", "county": "montco"}).json()
    pid = project["id"]
    upload = client.post(f"/api/projects/{pid}/imports", files={"file": (MONTCO_PDF.name, MONTCO_PDF.read_bytes(), "application/pdf")})
    assert upload.status_code == 201
    assert run_jobs() == 1
    import_id = upload.json()["import"]["id"]
    sale_list = client.get(f"/api/imports/{import_id}").json()
    assert sale_list["status"] == "review" and sale_list["stats"]["records"] == 55
    assert sale_list["detection"]["pages"][0]["kind"] == "notice"
    assert client.get(f"/api/imports/{import_id}/pages/2").headers["content-type"] == "image/png"

    rows = client.get(f"/api/imports/{import_id}/rows").json()["rows"]
    row = rows[0]
    assert row["page_number"] == 2 and row["values"]["parcel"]["value"] == "01-00-01606-02-2"
    corrected = client.post(f"/api/rows/{row['id']}/corrections", json={"field": "owner_name", "value": "REVIEWED OWNER LLC"}).json()
    assert corrected["values"]["owner_name"]["corrected"] is True
    assert client.post(f"/api/imports/{import_id}/parse", json={}).status_code == 409
    client.patch(f"/api/rows/{rows[1]['id']}", json={"excluded": True})
    committed = client.post(f"/api/imports/{import_id}/commit", json={}).json()
    assert committed["created"] == 54

    pilot = client.post(f"/api/projects/{pid}/pilot", json={"count": 8}).json()["pilot"]
    assert len(pilot) == 8
    sources = client.get("/api/sources", params={"county": "montco"}).json()
    assert next(s for s in sources if s["key"] == "montco.tax_claim")["requires_terms_ack"] is True
    ack = client.patch("/api/sources/montco.tax_claim", json={"acknowledge_terms": True, "note": "personal research"}).json()
    assert ack["terms_acknowledged_by"] == "local-user"

    with respx.mock(assert_all_called=False) as mock:
        gis_router(mock)
        taxclaim_router(mock, resolved_parcels={"02-00-02904-00-1"})
        job_id = client.post(f"/api/projects/{pid}/enrich", json={"selection": "pilot"}).json()["job_id"]
        processed = run_jobs()
    assert processed == 9
    job = client.get(f"/api/jobs/{job_id}").json()
    assert job["status"] == "succeeded" and job["children"]["succeeded"] == 8

    grid = client.get(f"/api/projects/{pid}/grid").json()
    assert len(grid["rows"]) == 54 and grid["columns"][0]["header"] == "Municipality"
    by_parcel = {r["values"]["parcel"]: r for r in grid["rows"]}
    resolved = by_parcel["02-00-02904-00-1"]
    assert resolved["styles"]["row"] == "red" and resolved["values"]["decision_status"] == "Removed/resolved"
    duplex = by_parcel["02-00-01016-00-8"]
    assert duplex["values"]["land_use_description"] == "R - DUPLEX" and duplex["quality"]["land_use_description"] == "verified"
    assert duplex["values"]["owner_name"] and duplex["sources"]["land_use_description"].startswith("Montgomery County GIS")

    tasks = client.get(f"/api/projects/{pid}/tasks").json()
    kinds = [t["source_key"] for t in tasks]
    assert kinds.count("montco.recorder") == 8 and kinds.count("montco.civil") == 8
    # Pilot parcels without a GIS record fall back to a user-assisted portal task instead of invented values.
    no_gis = [r for r in grid["rows"] if r["is_pilot"] and r["values"]["land_use_description"] is None]
    assert kinds.count("montco.assessment_portal") == len(no_gis)

    prop_id = duplex["id"]
    recorder_html = (FIXTURES / "capture" / "recorder_results_synthetic.html").read_text(encoding="utf-8")
    preview = client.post(f"/api/properties/{prop_id}/capture", json={"source_key": "montco.recorder", "content": recorder_html, "preview": True}).json()
    assert preview["status"] == "success" and preview["saved"] is False and len(preview["payload"]) == 10
    saved = client.post(f"/api/properties/{prop_id}/capture", json={"source_key": "montco.recorder", "content": recorder_html}).json()
    assert saved["saved"] is True
    civil = client.post(f"/api/properties/{prop_id}/capture/file", data={"source_key": "montco.civil"},
                        files={"file": ("results.html", (FIXTURES / "capture" / "montco_civil_results_synthetic.html").read_bytes(), "text/html")}).json()
    assert civil["saved"] is True
    bad = client.post(f"/api/properties/{prop_id}/capture", json={"source_key": "delco.civil", "content": "x"})
    assert bad.status_code == 400

    detail = client.get(f"/api/properties/{prop_id}").json()
    statuses = {m["effective_status"] for m in detail["mortgages"]}
    assert "No satisfaction found" in statuses and "Satisfied" in statuses
    assert detail["cells"]["mortgage_summary"]["value"].startswith("4 MTG(s)")  # the 1987 mortgage predates 1990
    assert len(detail["liens"]) == 1 and detail["cells"]["land_use_description"]["raw"] == "1132"
    review = next(m for m in detail["mortgages"] if m["effective_status"].startswith("Potential"))
    confirmed = client.post(f"/api/mortgages/{review['id']}/matches",
                            json={"satisfaction_document_id": review["matches"][0]["satisfaction_document_id"], "confirm": True, "note": "Servicer release"}).json()
    assert confirmed["status"] == "Satisfied"

    manual = client.post(f"/api/properties/{prop_id}/valuations/manual", json={"site": "zillow", "value": 315400, "note": "viewed by analyst"})
    assert manual.status_code == 201
    override = client.post(f"/api/properties/{prop_id}/overrides", json={"field": "living_area_sqft", "value": "1800", "reason": "Measured on site visit"})
    assert override.status_code == 201
    assert client.post(f"/api/properties/{prop_id}/overrides", json={"field": "decision_status", "value": "Maybe", "reason": "xx x"}).status_code == 400
    cell = client.get(f"/api/properties/{prop_id}").json()["cells"]
    assert cell["living_area_sqft"]["value"] == 1800 and cell["living_area_sqft"]["quality"] == "manual"
    assert cell["selected_valuation"]["value"] == 315400 and "Lowest credible" in cell["valuation_detail"]["value"]

    rules = client.get("/api/rules").json()
    assert rules["active"]["config"]["valuation"]["interest_threshold"] == 200000
    updated = client.put("/api/rules", json={"patch": {"valuation": {"interest_threshold": 350000}}, "reason": "Raise bar for 2026"})
    assert updated.status_code == 200 and updated.json()["version"] == 2
    assert run_jobs() == 1
    grid = client.get(f"/api/projects/{pid}/grid").json()
    row = next(r for r in grid["rows"] if r["id"] == prop_id)
    assert row["values"]["decision_status"] == "Not interested" and row["styles"]["cells"]["selected_valuation"] == "orange"

    providers = client.get("/api/providers").json()
    assert {p["key"] for p in providers["valuation"]} >= {"attom", "rentcast", "zillow_bridge", "sandbox"}
    assert client.put("/api/providers/attom", json={"enabled": True}).status_code == 400
    secret = client.put("/api/secrets/provider.attom.api_key", json={"value": "attom-key-abcdef-1234"}).json()
    assert secret["configured"] and "attom-key" not in str(secret)
    assert client.put("/api/secrets/not.allowed", json={"value": "whatever"}).status_code == 400

    export = client.post(f"/api/projects/{pid}/exports?sync=true", json={}).json()
    download = client.get(export["download"])
    assert download.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(download.content))
    assert wb["Updated Sale List"].max_row == 55 and "Lien Cases" in wb.sheetnames

    audit = client.get("/api/audit", params={"project_id": pid}).json()
    actions = {a["action"] for a in audit}
    assert {"import.upload", "import.correct_value", "import.commit", "enrichment.lookup", "property.override", "export.xlsx"} <= actions
    assert client.get("/api/audit/verify").json()["intact"] is True
    assert "attom-key-abcdef-1234" not in str(client.get("/api/audit").json())


def test_openapi_documents_endpoints(client):
    paths = client.get("/openapi.json").json()["paths"]
    for path in ("/api/projects", "/api/projects/{project_id}/imports", "/api/imports/{import_id}/commit", "/api/projects/{project_id}/enrich",
                 "/api/properties/{property_id}/capture", "/api/projects/{project_id}/refresh-tax-status", "/api/projects/{project_id}/exports"):
        assert path in paths


def test_token_auth_mode(client, monkeypatch):
    from app.api.deps import hash_token
    from app.config import get_settings
    from app.db import session_scope
    from app.models import User

    monkeypatch.setattr(get_settings(), "auth_mode", "token")
    assert client.get("/api/projects").status_code == 401
    with session_scope() as s:
        s.add(User(email="viewer@example.com", display_name="V", role="viewer", token_hash=hash_token("usi_viewer_token")))
    headers = {"Authorization": "Bearer usi_viewer_token"}
    assert client.get("/api/projects", headers=headers).status_code == 200
    assert client.post("/api/projects", json={"name": "x", "county": "delco"}, headers=headers).status_code == 403
