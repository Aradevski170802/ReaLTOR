import openpyxl
import pytest
import respx
from sqlalchemy import select

from app.enums import Decision, MortgageStatus
from app.export.template_inspector import load_builtin
from app.export.xlsx_exporter import build_workbook, export_project
from app.jobs.worker import Worker
from app.models import CaptureTask, MortgageRecord, Project, Property, RefreshRun, TaxStatusHistory, UserOverride
from app.seed import seed_demo
from app.services.imports import OverwriteConflict, commit_import, correct_value, run_parse, save_upload
from app.services.policies import acknowledge_terms
from app.services.refresh import refresh_tax_status
from app.services.rules_service import active_config, evaluate_properties
from app.services.snapshot import load_snapshots
from tests.conftest import DELCO_PDF, MONTCO_PDF, gis_router, taxclaim_router


def make_project(db, county: str, pdf) -> Project:
    project = Project(name=f"Test {county}", county=county, created_by="tester")
    db.add(project)
    db.flush()
    sale_list = save_upload(db, project, pdf.name, pdf.read_bytes(), "tester")
    run_parse(db, sale_list, "tester")
    commit_import(db, sale_list, "tester")
    db.commit()
    return project


def test_import_corrections_are_protected(db):
    project = Project(name="P", county="delco")
    db.add(project)
    db.flush()
    sale_list = save_upload(db, project, DELCO_PDF.name, DELCO_PDF.read_bytes(), "tester")
    run_parse(db, sale_list, "tester")
    row = sale_list.rows[0]
    correct_value(db, row, "owner_name", "CORRECTED OWNER", "reviewer")
    with pytest.raises(OverwriteConflict):
        run_parse(db, sale_list, "tester")
    result = commit_import(db, sale_list, "tester")
    assert result.created == 15 and not result.skipped
    prop = db.scalar(select(Property).where(Property.project_id == project.id, Property.sale_number == "1"))
    assert prop.owner_name == "CORRECTED OWNER" and prop.source_page == 1
    assert {i.kind for i in prop.identifiers} >= {"folio", "gis_parid", "datalet_pin"}
    assert [t.tax_year for t in db.get(Property, prop.id).tax_lines] == [2024]
    db.add(UserOverride(property_id=prop.id, field="owner_name", override_value="MY OVERRIDE", reason="verified deed", created_by="x"))
    correct_value(db, row, "owner_name", "ANOTHER NAME", "reviewer")
    conflict = commit_import(db, sale_list, "tester")
    assert conflict.conflicts and conflict.conflicts[0]["fields"] == ["owner_name"]
    assert commit_import(db, sale_list, "tester", confirm_overwrite=True).updated >= 1


def test_montco_enrichment_flow_and_refresh(db):
    project = make_project(db, "montco", MONTCO_PDF)
    props = db.scalars(select(Property).where(Property.project_id == project.id).order_by(Property.sequence, Property.id)).all()
    pilot = props[:6]
    from app.services.enrichment import enrich_property

    with respx.mock(assert_all_called=False) as mock:
        gis_router(mock)
        taxclaim_router(mock, resolved_parcels={"02-00-02904-00-1"})
        summary = enrich_property(db, pilot[0])
        assert summary["montco.assessment_gis"] == "success"
        assert summary["montco.tax_claim"] == "terms_not_acknowledged"
        acknowledge_terms(db, "montco.tax_claim", "admin", "Personal research use confirmed")
        for prop in pilot:
            enrich_property(db, prop, force=True)
        db.commit()
        config = active_config(db)
        snaps = load_snapshots(db, [p.id for p in pilot], config)
        brandi = next(s for s in snaps.values() if s.prop.parcel_normalized == "02-00-02904-00-1")
        assert brandi.value("tax_2024_balance") == 0.0 and brandi.styles()["row"] == "red"
        assert brandi.prop.decision_status == Decision.REMOVED
        duplex = next(s for s in snaps.values() if s.prop.parcel_normalized == "02-00-01016-00-8")
        assert duplex.value("land_use_description") == "R - DUPLEX" and duplex.value("living_area_sqft") == 1724
        cell = duplex.cell("land_use_description")
        assert cell.raw == "1132" and cell.quality == "verified" and cell.url.startswith("https://gis.montcopa.org")
        # The county GIS record has a blank LOC_ZIP for this parcel, so no ZIP is invented and the value is marked estimated.
        assert duplex.value("full_location") == "1002 DEKALB ST, Bridgeport, PA"
        assert duplex.cell("full_location").quality == "estimated"
        assert duplex.value("selected_valuation") is None and duplex.prop.decision_status == Decision.INSUFFICIENT
        tasks = db.scalars(select(CaptureTask).where(CaptureTask.property_id == duplex.prop.id)).all()
        assert {t.source_key for t in tasks} == {"montco.recorder", "montco.civil"}

        run = refresh_tax_status(db, project, "manual", property_ids=[p.id for p in pilot])
        db.commit()
    assert run.total == 6 and run.updated == 6 and run.failed == 0
    assert db.scalar(select(RefreshRun).where(RefreshRun.id == run.id)).summary.startswith("6 checked")
    history = db.scalars(select(TaxStatusHistory).where(TaxStatusHistory.refresh_run_id == run.id)).all()
    assert len(history) == 6


def test_delco_refresh_creates_recapture_tasks(db):
    project = make_project(db, "delco", DELCO_PDF)
    run = refresh_tax_status(db, project, "scheduled")
    db.commit()
    assert run.user_action_needed == run.total == 15
    tasks = db.scalars(select(CaptureTask).where(CaptureTask.project_id == project.id, CaptureTask.reason == "daily_refresh")).all()
    assert len(tasks) == 15 and tasks[0].url.startswith("http://delcorealestate")


def test_seed_demo_and_export_workbook(db):
    result = seed_demo()
    assert result["montco"]["created"] and result["delco"]["created"]
    assert seed_demo()["montco"]["created"] is False
    config = active_config(db)

    delco = db.get(Project, result["delco"]["project_id"])
    d_ids = result["delco"]["pilot"]
    d_snaps = load_snapshots(db, d_ids, config)
    styles = [s.styles()["cells"].get("tax_sale_status") for s in d_snaps.values()]
    assert styles.count("green") == 6 and styles.count("red") == 2
    first = d_snaps[d_ids[0]]
    assert first.value("selected_valuation") is not None and first.cell("selected_valuation").quality == "estimated"
    assert first.value("mortgage_summary") == "4 MTG(s)- 2 SAT, 1 NO SAT, 1 REVIEW"
    assert first.cell("property_type").notes.startswith("DEMO FIXTURE")
    statuses = {m.effective_status for m in db.scalars(select(MortgageRecord).where(MortgageRecord.property_id == d_ids[0]))}
    assert statuses == {MortgageStatus.SATISFIED, MortgageStatus.NO_SATISFACTION, MortgageStatus.REVIEW}

    montco = db.get(Project, result["montco"]["project_id"])
    wb, count = build_workbook(db, montco)
    assert count == 55
    assert wb.sheetnames[:2] == ["Updated Sale List", "Mortgage Info"]
    for name in ("Lien Cases", "Sources", "Source Catalog", "Refresh Log", "Tax Status History", "Errors & Review Queue",
                 "User Action Tasks", "PDF Extraction", "Configuration", "Audit Log", "Notes & Legend"):
        assert name in wb.sheetnames
    ws = wb["Updated Sale List"]
    reference = [c["header"] for c in load_builtin("montco")["sheets"][0]["columns"]]
    headers = [ws.cell(1, i).value for i in range(1, 37)]
    assert headers[:28] == reference[:28] and headers[29:] == reference[29:]
    assert headers[28] in ("Selected Market Valuation", "Lowest Across Sites")
    assert ws.freeze_panes == "D2" and ws.auto_filter.ref.startswith("A1:")
    assert ws.cell(1, 1).fill.fgColor.rgb == "FF01696F" and ws.column_dimensions["E"].width == 34
    formulas = [rule.formula[0] for rng in ws.conditional_formatting for rule in rng.rules]
    assert any("ROUND($AH2,2)=0" in f for f in formulas) and any('SEARCH("NO SAT"' in f for f in formulas)
    parcel_col = headers.index("Parcel") + 1
    price_col = headers.index("Approx. Sale Price") + 1
    assert isinstance(ws.cell(2, price_col).value, float) and ws.cell(2, price_col).number_format.startswith('"$"')
    assert ws.cell(2, parcel_col).value
    mortgage = wb["Mortgage Info"]
    assert [mortgage.cell(1, i).value for i in range(1, 12)] == [c["header"] for c in load_builtin("montco")["sheets"][1]["columns"]]
    assert any(mortgage.cell(r, 11).value == "NO SATISFACTION" for r in range(2, mortgage.max_row + 1))
    assert wb["Sources"].max_row > 50 and wb["PDF Extraction"].max_row == 56

    delco_wb, delco_count = build_workbook(db, delco)
    dws = delco_wb["Updated Sale List"]
    d_headers = [dws.cell(1, i).value for i in range(1, dws.max_column + 1)]
    assert delco_count == 15 and d_headers[2] == "Folio-NBR" and "Tax Sale Status" in d_headers
    dformulas = [rule.formula[0] for rng in dws.conditional_formatting for rule in rng.rules]
    assert any('"Listed for Upset Sale"' in f for f in dformulas) and any("<100000" in f for f in dformulas)

    record = export_project(db, montco, "tester")
    db.commit()
    reopened = openpyxl.load_workbook(record.file_path)
    assert reopened["Updated Sale List"].max_row == 56 and record.sha256


def test_worker_processes_export_job(db):
    from app.jobs.queue import enqueue

    project = make_project(db, "delco", DELCO_PDF)
    evaluate_properties(db, [p.id for p in project.properties], "tester")
    job = enqueue(db, "export_workbook", {}, project_id=project.id, created_by="tester")
    db.commit()
    assert Worker(enable_scheduler=False).run_until_empty() == 1
    db.expire_all()
    assert db.get(type(job), job.id).status == "succeeded"


def test_autopilot_commits_and_enriches_montco(db):
    from app.jobs.queue import enqueue
    from app.services.policies import acknowledge_automated_terms

    project = Project(name="Autopilot montco", county="montco", created_by="tester")
    db.add(project)
    db.flush()
    sale_list = save_upload(db, project, MONTCO_PDF.name, MONTCO_PDF.read_bytes(), "tester")
    run_parse(db, sale_list, "tester")  # status = review, nothing committed yet
    acknowledge_automated_terms(db, "tester")  # so the Tax Claim source runs automatically
    assert db.scalar(select(Property).where(Property.project_id == project.id)) is None
    enqueue(db, "autopilot", {}, project_id=project.id, created_by="tester")
    db.commit()

    with respx.mock(assert_all_called=False) as mock:
        gis_router(mock)
        taxclaim_router(mock, resolved_parcels={"02-00-02904-00-1"})
        processed = Worker(enable_scheduler=False).run_until_empty()
    assert processed >= 55  # 1 autopilot + one enrich_property per committed row
    db.expire_all()
    props = db.scalars(select(Property).where(Property.project_id == project.id)).all()
    assert len(props) == 55
    assert all(p.decision_status for p in props)  # every property was scored with no manual step
    duplex = db.scalar(select(Property).where(Property.parcel_normalized == "02-00-01016-00-8"))
    snap = load_snapshots(db, [duplex.id], active_config(db))[duplex.id]
    assert snap.value("land_use_description") == "R - DUPLEX"  # GIS ran automatically
    assert snap.value("tax_2024_balance") is not None  # Tax Claim ran automatically
