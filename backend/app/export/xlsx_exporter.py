"""Research workbook export that follows the reference workbook layout and adds audit sheets."""

from __future__ import annotations

import hashlib
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.orm import Session

from app import __version__
from app.audit import record_audit
from app.config import get_settings
from app.connectors.registry import get_connector
from app.db import utcnow
from app.enums import COUNTY_LABELS, MortgageStatus
from app.export.template_inspector import header_style_lookup, load_builtin, sheet_spec
from app.export.workbook_spec import (
    ADDITION_HEADER_FILL,
    REFERENCE_HEADER_FILL,
    ColumnSpec,
    column_letter,
    main_columns,
)
from app.models import (
    AuditLog,
    CaptureTask,
    CivilLienCase,
    ExportRecord,
    ExtractedRow,
    FieldProvenance,
    ImportedSaleList,
    Project,
    Property,
    ProviderConfig,
    RecorderDocument,
    RefreshRun,
    SourceEvidence,
    SourceLookup,
    SourcePolicy,
    TaxStatusHistory,
    WorkbookTemplate,
)
from app.services.rules_service import active_ruleset
from app.services.snapshot import PropertySnapshot, load_snapshots

RED, GREEN, ORANGE, YELLOW, GREY = "FFFF0000", "FF00B050", "FFFFC000", "FFFFEB9C", "FFD9D9D9"
WHITE, DARK = "FFFFFFFF", "FF000000"
DISCLAIMER = (
    "Informational research only — not legal, title, appraisal, tax, investment or lien-clearance advice. County data "
    "sources are not certified searches; automated valuations are estimates. Verify everything with official records "
    "and qualified professionals before bidding."
)
_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _fill(rgb: str) -> PatternFill:
    return PatternFill(fill_type="solid", start_color=rgb, end_color=rgb, fgColor=rgb)


def _clean(value: Any) -> Any:
    if isinstance(value, str):
        return _ILLEGAL.sub("", value)[:32000]
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    return value


def _excel_value(value: Any, kind: str) -> Any:
    if value is None or value == "":
        return None
    if kind in ("money", "number"):
        try:
            return float(value)
        except (TypeError, ValueError):
            return _clean(value)
    if kind == "int":
        try:
            return int(value)
        except (TypeError, ValueError):
            return _clean(value)
    if kind == "date" and isinstance(value, str):
        try:
            return date.fromisoformat(value[:10])
        except ValueError:
            return value
    return _clean(value)


def _header(ws: Worksheet, headers: list[str], fill: str = "FF424242", widths: list[float] | None = None, height: float = 30) -> None:
    for idx, text in enumerate(headers, start=1):
        cell = ws.cell(1, idx, text)
        cell.font = Font(name="Calibri", size=11, bold=True, color=WHITE)
        cell.fill = _fill(fill)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if widths:
            ws.column_dimensions[get_column_letter(idx)].width = widths[idx - 1]
    ws.row_dimensions[1].height = height
    ws.freeze_panes = "A2"


def _finish_table(ws: Worksheet, ncols: int) -> None:
    if ws.max_row >= 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}{max(ws.max_row, 2)}"


def _absolute(url: str | None) -> str | None:
    if not url:
        return None
    if url.startswith("/"):
        base = getattr(get_settings(), "public_base_url", None) or "http://localhost:8000"
        return base.rstrip("/") + url
    return url


def _template_for(session: Session, county: str) -> dict | None:
    uploaded = session.scalar(select(WorkbookTemplate).where(WorkbookTemplate.county == county).order_by(WorkbookTemplate.id.desc()))
    return uploaded.spec if uploaded else load_builtin(county) or load_builtin("montco")


# ------------------------------------------------------------------------------------------------ main sheet


def write_main_sheet(ws: Worksheet, snaps: list[PropertySnapshot], columns: list[ColumnSpec], config: dict, county: str,
                     template: dict | None) -> None:
    ws.title = "Updated Sale List"
    styles = header_style_lookup(template, "Updated Sale List")
    tsheet = sheet_spec(template, "Updated Sale List") or {}
    header_height = tsheet.get("header_row_height") or 36
    data_height = tsheet.get("data_row_height") or 45

    for idx, col in enumerate(columns, start=1):
        letter = get_column_letter(idx)
        tcol = styles.get(col.header.strip().lower()) if col.reference else None
        cell = ws.cell(1, idx, col.header)
        fill = (tcol or {}).get("header_fill") or (REFERENCE_HEADER_FILL if col.reference else ADDITION_HEADER_FILL)
        font = (tcol or {}).get("header_font") or {}
        cell.fill = _fill(fill)
        cell.font = Font(name=font.get("name") or "Calibri", size=font.get("size") or col.header_font_size, bold=True,
                         color=font.get("color") or WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[letter].width = (tcol or {}).get("width") or col.width
        if col.hidden:
            ws.column_dimensions[letter].hidden = True
    ws.row_dimensions[1].height = header_height

    thin = Side(style="thin", color="FFBFBFBF")
    for r, snap in enumerate(snaps, start=2):
        ws.row_dimensions[r].height = data_height
        for idx, col in enumerate(columns, start=1):
            if col.key == "tax_status_override":
                ov = snap.overrides.get("tax_status_interpretation")
                value = f"Overridden: {ov.reason}" if ov else None
                cell_meta = None
            else:
                cell_meta = snap.cell(col.key)
                value = cell_meta.value
            xcell = ws.cell(r, idx, _excel_value(value, col.kind))
            xcell.alignment = Alignment(wrap_text=True, vertical="top")
            xcell.border = Border(bottom=thin)
            if col.number_format and not isinstance(xcell.value, str):
                xcell.number_format = col.number_format
            if cell_meta is not None:
                if cell_meta.quality == "estimated":
                    xcell.font = Font(italic=True)
                elif cell_meta.quality == "manual":
                    xcell.font = Font(color="FF1F4E79")
                elif cell_meta.quality == "stale":
                    xcell.font = Font(color="FF9C5700", italic=True)
            if col.key == "parcel":
                link = next((pi.value for pi in snap.prop.identifiers if pi.kind == "detail_link"), None)
                link = link or (snap.assessments.get("montco.assessment_gis").detail_link if snap.assessments.get("montco.assessment_gis") else None)
                if link:
                    xcell.hyperlink = link
                    xcell.font = Font(color="FF0563C1", underline="single")
            if col.key == "property_photo" and isinstance(value, str) and value.startswith("/api/"):
                xcell.value = "View photo"
                xcell.hyperlink = _absolute(value)
                xcell.font = Font(color="FF0563C1", underline="single")

    last_row = max(2, len(snaps) + 1)
    last_col = get_column_letter(len(columns))
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{last_col}{last_row}"
    _conditional_formats(ws, columns, config, county, last_row)


def _conditional_formats(ws: Worksheet, columns: list[ColumnSpec], config: dict, county: str, last_row: int) -> None:
    pos = {c.key: column_letter(i) for i, c in enumerate(columns)}
    full = f"A2:{get_column_letter(len(columns))}{last_row}"
    white_bold = Font(color=WHITE, bold=True)
    threshold = config["valuation"]["interest_threshold"]

    def add(ref: str, formula: str, fill: str, font: Font | None = None, stop: bool = False) -> None:
        ws.conditional_formatting.add(ref, FormulaRule(formula=[formula], fill=_fill(fill), font=font, stopIfTrue=stop))

    if county == "montco" and "tax_2024_balance" in pos:
        bal, ov = pos["tax_2024_balance"], pos.get("tax_status_override")
        cond = f"AND(ISNUMBER(${bal}2),ROUND(${bal}2,2)=0" + (f",LEN(${ov}2)=0" if ov else "") + ")"
        if config["tax"]["montco_zero_2024_balance_resolved"]:
            add(full, cond, RED, white_bold, stop=True)
    if county == "delco":
        if "tax_sale_status" in pos:
            s = pos["tax_sale_status"]
            add(f"{s}2:{s}{last_row}", f'TRIM(${s}2)="Listed for Upset Sale"', GREEN, white_bold)
            add(f"{s}2:{s}{last_row}", f'TRIM(${s}2)<>"Listed for Upset Sale"', RED, white_bold)
        if "assessed_value" in pos:
            a = pos["assessed_value"]
            add(f"{a}2:{a}{last_row}", f"AND(ISNUMBER(${a}2),${a}2<{config['assessment']['highlight_below']})", ORANGE)
    for key in ("val_homes_com", "val_zillow", "val_realtor", "val_redfin", "val_trulia", "selected_valuation", "val_attom", "val_rentcast", "val_sandbox"):
        if key in pos:
            L = pos[key]
            add(f"{L}2:{L}{last_row}", f"AND(ISNUMBER(${L}2),${L}2<{threshold})", ORANGE)
    if "market_value_flag" in pos:
        L = pos["market_value_flag"]
        add(f"{L}2:{L}{last_row}", f'LEFT(${L}2,5)="under"', ORANGE)
    if "mortgage_summary" in pos:
        L = pos["mortgage_summary"]
        rng = f"{L}2:{L}{last_row}"
        add(rng, f'ISNUMBER(SEARCH("NO SAT",${L}2))', RED, white_bold, stop=True)
        add(rng, f'ISNUMBER(SEARCH("REVIEW",${L}2))', YELLOW, stop=True)
        add(rng, f'ISNUMBER(SEARCH("OK(SATs",${L}2))', GREEN, white_bold)
    if "decision_status" in pos:
        L = pos["decision_status"]
        rng = f"{L}2:{L}{last_row}"
        for text, fill, font in (("Interested", GREEN, white_bold), ("Not interested", ORANGE, None), ("Review required", YELLOW, None),
                                 ("Insufficient data", GREY, None), ("Removed/resolved", RED, white_bold)):
            add(rng, f'${L}2="{text}"', fill, font)


# ------------------------------------------------------------------------------------------------ supporting sheets


def write_mortgage_sheet(ws: Worksheet, session: Session, snaps: list[PropertySnapshot], config: dict) -> None:
    headers = ["Parcel", "Municipality", "Instrument", "Type", "Recorded Date", "Grantor", "Grantee", "Mortgage Amt", "Satisfied",
               "Assoc Instruments", "Document Status", "Match Status", "Matched Satisfaction(s)", "Match Reasons", "Source",
               "Source URL", "Entry Method"]
    widths = [16, 14, 18, 28, 13, 28, 28, 14, 10, 22, 16, 26, 22, 50, 26, 30, 12]
    _header(ws, headers, "FF424242", widths)
    since = date(config["mortgage"]["since_year"], 1, 1)
    status_text = {MortgageStatus.SATISFIED: "SATISFIED", MortgageStatus.NO_SATISFACTION: "NO SATISFACTION",
                   MortgageStatus.REVIEW: "REVIEW REQUIRED", MortgageStatus.INSUFFICIENT: "INSUFFICIENT DATA"}
    row = 2
    for snap in snaps:
        docs = session.scalars(select(RecorderDocument).where(RecorderDocument.property_id == snap.prop.id)
                               .order_by(RecorderDocument.recorded_date.desc().nullslast())).all()
        mortgages = {m.document_id: m for m in snap.mortgages}
        doc_by_id = {d.id: d for d in docs}
        for d in docs:
            m = mortgages.get(d.id)
            instrument = d.instrument_number or (f"{d.book} {d.page}" if d.book and d.page else None)
            satisfied = doc_status = match_status = matched = reasons = None
            if d.doc_category == "mortgage":
                if m is not None:
                    eff = m.effective_status
                    doc_status = status_text.get(eff, eff)
                    match_status = eff + (" (reviewer)" if m.user_status else "")
                    satisfied = {"Satisfied": "Yes", "No satisfaction found": "No"}.get(eff, "Review")
                    confirmed = [x for x in m.matches if x.decision in ("auto_confirmed", "user_confirmed")]
                    candidates = confirmed or [x for x in m.matches if x.decision == "candidate"]
                    matched = "; ".join(
                        (doc_by_id[x.satisfaction_document_id].instrument_number or f"doc {x.satisfaction_document_id}") + f" ({x.decision}, {x.score:.2f})"
                        for x in candidates if x.satisfaction_document_id in doc_by_id) or None
                    reasons = m.status_reason
                elif d.recorded_date and d.recorded_date < since:
                    doc_status = f"PRE-{since.year} (not evaluated)"
            values = [snap.prop.parcel_normalized, snap.prop.municipality, instrument, d.doc_type_raw, d.recorded_date, d.grantors,
                      d.grantees, d.amount if d.doc_category == "mortgage" else None, satisfied, d.related_reference, doc_status,
                      match_status, matched, reasons, snap.lookups.get(d.source_key).source_key if snap.lookups.get(d.source_key) else d.source_key,
                      d.source_url, d.entry_method]
            for c, v in enumerate(values, start=1):
                cell = ws.cell(row, c, _clean(v))
                if c == 5 and isinstance(v, date):
                    cell.number_format = "mm/dd/yyyy"
                if c == 8 and v is not None:
                    cell.number_format = '"$"#,##0.00'
            if d.source_url:
                ws.cell(row, 16).hyperlink = d.source_url
            row += 1
    last = max(row - 1, 2)
    rng = f"A2:{get_column_letter(len(headers))}{last}"
    white_bold = Font(color=WHITE, bold=True)
    ws.conditional_formatting.add(rng, FormulaRule(formula=['$K2="NO SATISFACTION"'], fill=_fill(RED), font=white_bold))
    ws.conditional_formatting.add(rng, FormulaRule(formula=['$K2="SATISFIED"'], fill=_fill(GREEN), font=white_bold))
    ws.conditional_formatting.add(rng, FormulaRule(formula=['OR($K2="REVIEW REQUIRED",$K2="INSUFFICIENT DATA")'], fill=_fill(YELLOW)))
    _finish_table(ws, len(headers))


def write_simple(ws: Worksheet, headers: list[str], widths: list[float], rows: list[list[Any]], link_cols: tuple[int, ...] = ()) -> None:
    _header(ws, headers, "FF424242", widths)
    for r, values in enumerate(rows, start=2):
        for c, v in enumerate(values, start=1):
            cell = ws.cell(r, c, _clean(v))
            if isinstance(v, datetime):
                cell.number_format = "mm/dd/yyyy hh:mm"
            elif isinstance(v, date):
                cell.number_format = "mm/dd/yyyy"
            if c in link_cols and isinstance(v, str) and v.startswith("http"):
                cell.hyperlink = v
                cell.font = Font(color="FF0563C1", underline="single")
    _finish_table(ws, len(headers))


def build_workbook(session: Session, project: Project, property_ids: list[int] | None = None) -> tuple[Workbook, int]:
    ruleset = active_ruleset(session)
    config = ruleset.config
    ids_query = select(Property.id).where(Property.project_id == project.id, Property.excluded.is_(False))
    if property_ids:
        ids_query = ids_query.where(Property.id.in_(property_ids))
    ids = list(session.scalars(ids_query))
    snap_map = load_snapshots(session, ids, config)
    snaps = sorted(snap_map.values(), key=lambda s: ((s.prop.municipality or "").upper(), s.prop.sequence or 0, s.prop.id))
    columns = main_columns(project.county, config, project.is_demo)
    template = _template_for(session, project.county)
    connector = get_connector(project.county)

    wb = Workbook()
    write_main_sheet(wb.active, snaps, columns, config, project.county, template)
    write_mortgage_sheet(wb.create_sheet("Mortgage Info"), session, snaps, config)

    prop_by_id = {s.prop.id: s.prop for s in snaps}
    liens = session.scalars(select(CivilLienCase).where(CivilLienCase.property_id.in_(ids or [0]))).all()
    write_simple(wb.create_sheet("Lien Cases"),
                 ["Parcel", "Owner", "Source", "Search Term", "Case Number", "Status", "Classification", "Case Type", "Parties",
                  "Filing Date", "Amount", "Why Retained", "Source URL", "Retrieved", "Entry Method"],
                 [16, 24, 22, 22, 16, 14, 22, 18, 36, 12, 12, 30, 30, 16, 12],
                 [[prop_by_id[c.property_id].parcel_normalized, prop_by_id[c.property_id].owner_name, c.source_key, c.search_term,
                   c.case_number, c.status, c.classification, c.case_type, c.parties, c.filing_date, c.amount, c.relevance_reason,
                   c.source_url, c.retrieved_at, c.entry_method] for c in liens if c.is_relevant], link_cols=(13,))

    provenance = session.scalars(select(FieldProvenance).where(FieldProvenance.property_id.in_(ids or [0]), FieldProvenance.is_current.is_(True))
                                 .order_by(FieldProvenance.property_id, FieldProvenance.field)).all()
    evidence = {e.id: e for e in session.scalars(select(SourceEvidence).where(SourceEvidence.property_id.in_(ids or [0])))}
    write_simple(wb.create_sheet("Sources"),
                 ["Parcel", "Field", "Value", "Raw Source Value", "Source", "Source URL", "Retrieved At", "Quality", "Confidence",
                  "Notes", "Evidence ID", "Evidence SHA-256", "Entry Method"],
                 [16, 22, 26, 20, 34, 34, 16, 10, 10, 40, 10, 20, 12],
                 [[prop_by_id[p.property_id].parcel_normalized, p.field, p.value, p.raw_value, p.source_name, p.source_url, p.retrieved_at,
                   p.quality, p.confidence, p.notes, p.evidence_id, evidence[p.evidence_id].sha256 if p.evidence_id in evidence else None,
                   evidence[p.evidence_id].entry_method if p.evidence_id in evidence else None] for p in provenance if p.property_id in prop_by_id],
                 link_cols=(6,))

    policies = {p.source_key: p for p in session.scalars(select(SourcePolicy))}
    write_simple(wb.create_sheet("Source Catalog"),
                 ["Source Key", "Name", "Category", "Access Method", "Compliance Status", "Terms / Access Notes", "Access Findings",
                  "Base URLs", "Rate Limit / min", "Cache TTL (h)", "Daily Refresh", "Enabled", "Terms Acknowledged By", "Acknowledged At", "Verified On"],
                 [26, 40, 12, 18, 22, 60, 50, 40, 10, 10, 10, 10, 18, 16, 12],
                 [[a.descriptor.key, a.descriptor.name, a.descriptor.category, str(a.descriptor.access_method), a.descriptor.compliance_status,
                   a.descriptor.terms_notes, " | ".join(a.descriptor.access_findings), " ".join(a.descriptor.base_urls),
                   a.descriptor.rate_limit_per_minute, a.descriptor.cache_ttl_hours, a.descriptor.daily_refresh,
                   policies[a.descriptor.key].enabled if a.descriptor.key in policies else True,
                   policies[a.descriptor.key].terms_acknowledged_by if a.descriptor.key in policies else None,
                   policies[a.descriptor.key].terms_acknowledged_at if a.descriptor.key in policies else None,
                   a.descriptor.verified_on] for a in connector.sources()])

    runs = session.scalars(select(RefreshRun).where(RefreshRun.project_id == project.id).order_by(RefreshRun.started_at.desc())).all()
    write_simple(wb.create_sheet("Refresh Log"),
                 ["Run ID", "Kind", "Trigger", "Started", "Finished", "Checked", "Updated", "Changed", "Failed", "User Action Needed", "Summary"],
                 [8, 12, 10, 16, 16, 9, 9, 9, 8, 10, 60],
                 [[r.id, r.kind, r.trigger, r.started_at, r.finished_at, r.total, r.updated, r.changed, r.failed, r.user_action_needed, r.summary] for r in runs])
    history = session.scalars(select(TaxStatusHistory).where(TaxStatusHistory.property_id.in_(ids or [0])).order_by(TaxStatusHistory.checked_at.desc())).all()
    write_simple(wb.create_sheet("Tax Status History"),
                 ["Parcel", "Checked At", "Source", "Lookup Status", "Status", "Status Category", "2024 Balance", "Delinquent Total", "Changed", "Refresh Run"],
                 [16, 16, 22, 18, 30, 14, 12, 14, 9, 10],
                 [[prop_by_id[h.property_id].parcel_normalized, h.checked_at, h.source_key, h.lookup_status, h.status_text, h.status_category,
                   h.balance_2024, h.delinquent_total, h.changed, h.refresh_run_id] for h in history if h.property_id in prop_by_id])

    queue_rows: list[list[Any]] = []
    for snap in snaps:
        p = snap.prop
        if p.needs_review:
            queue_rows.append([p.parcel_normalized, p.owner_name, "Extraction review", "sale_list", "needs_review",
                               "Sale-list row flagged during PDF extraction", p.created_at])
        for lk in snap.lookups.values():
            if lk.status not in ("success", "no_match", "skipped"):
                queue_rows.append([p.parcel_normalized, p.owner_name, "Source lookup", lk.source_key, lk.status, lk.message, lk.finished_at])
        for m in snap.mortgages:
            if m.effective_status in (MortgageStatus.REVIEW, MortgageStatus.INSUFFICIENT):
                queue_rows.append([p.parcel_normalized, p.owner_name, "Mortgage match review", m.document.source_key if m.document else None,
                                   m.effective_status, m.status_reason, m.evaluated_at])
        if p.decision_status in ("Review required", "Insufficient data"):
            queue_rows.append([p.parcel_normalized, p.owner_name, "Decision", "rules", p.decision_status, p.decision_summary, p.decision_at])
    open_tasks = session.scalars(select(CaptureTask).where(CaptureTask.project_id == project.id, CaptureTask.status == "open")).all()
    write_simple(wb.create_sheet("Errors & Review Queue"),
                 ["Parcel", "Owner", "Issue Type", "Source", "Status", "Detail", "Since"], [16, 24, 20, 26, 22, 70, 16], queue_rows)
    write_simple(wb.create_sheet("User Action Tasks"),
                 ["Parcel", "Source", "Reason", "Official URL", "Search", "Instructions", "Created", "Last Error"],
                 [16, 26, 16, 40, 40, 70, 16, 30],
                 [[prop_by_id[t.property_id].parcel_normalized if t.property_id in prop_by_id else t.property_id, t.source_key, t.reason, t.url,
                   t.search_hint, t.instructions, t.created_at, t.last_error] for t in open_tasks], link_cols=(4,))

    import_ids = {s.prop.import_id for s in snaps if s.prop.import_id}
    extraction_rows: list[list[Any]] = []
    for sale_list in session.scalars(select(ImportedSaleList).where(ImportedSaleList.id.in_(import_ids or {0}))):
        for er in session.scalars(select(ExtractedRow).where(ExtractedRow.import_id == sale_list.id).order_by(ExtractedRow.row_index)):
            vals = {v.field: v for v in er.values}

            def g(name: str, _vals=vals) -> str | None:
                return _vals[name].value if name in _vals else None

            corrected = ", ".join(f"{v.field} (by {v.corrected_by})" for v in er.values if v.corrected)
            extraction_rows.append([sale_list.filename, er.row_index + 1, er.page_number, g("sale_number"),
                                    vals["parcel"].raw_value if "parcel" in vals else None, g("parcel"), g("owner_name"),
                                    g("property_address"), g("municipality"), g("amount_due"), g("assessment"), g("description"),
                                    g("lot_size"), g("mailing_address"), er.confidence, er.needs_review,
                                    "; ".join(er.review_reasons or []), corrected or None, er.excluded, er.raw_text])
    write_simple(wb.create_sheet("PDF Extraction"),
                 ["File", "Row", "PDF Page", connector.parcel_label.replace("Folio-NBR", "Record #") if project.county == "delco" else "Sale Number",
                  "Parcel (as printed)", "Parcel (normalized)", "Owner", "Location", "Municipality", "Amount Due", "Assessment", "Description",
                  "Lot Size", "Mailing Address", "Confidence", "Needs Review", "Review Reasons", "Corrected Fields", "Excluded", "Raw Text"],
                 [24, 6, 8, 10, 18, 18, 26, 24, 16, 12, 12, 18, 16, 26, 10, 9, 40, 20, 8, 60], extraction_rows)

    providers = {p.provider_key: p for p in session.scalars(select(ProviderConfig))}
    config_rows: list[list[Any]] = [
        ["Project", project.name], ["County", COUNTY_LABELS.get(project.county, project.county)], ["Exported at (UTC)", utcnow()],
        ["Application version", __version__], ["Rule set version", ruleset.version], ["Rule set change reason", ruleset.change_reason],
        ["Demo project (sandbox values allowed)", project.is_demo], ["Disclaimer", DISCLAIMER],
    ]
    for section, values in config.items():
        for key, value in values.items():
            config_rows.append([f"rules.{section}.{key}", ", ".join(map(str, value)) if isinstance(value, list) else value])
    for key, p in providers.items():
        config_rows.append([f"provider.{key}.enabled", p.enabled])
    write_simple(wb.create_sheet("Configuration"), ["Setting", "Value"], [44, 100], config_rows)

    audits = session.scalars(select(AuditLog).where(AuditLog.project_id == project.id).order_by(AuditLog.id.desc()).limit(5000)).all()
    write_simple(wb.create_sheet("Audit Log"), ["ID", "When (UTC)", "Actor", "Action", "Entity", "Entity ID", "Details", "Hash"],
                 [8, 16, 18, 24, 18, 10, 80, 20],
                 [[a.id, a.occurred_at, a.actor, a.action, a.entity_type, a.entity_id, str(a.details)[:3000], a.hash[:16]] for a in audits])

    legend = wb.create_sheet("Notes & Legend")
    legend.column_dimensions["A"].width = 110
    lines = [
        ("How to use this workbook", True),
        ("", False),
        (f"{COUNTY_LABELS.get(project.county)} — generated by Upset Sale Intel {__version__} on {utcnow():%Y-%m-%d %H:%M} UTC.", False),
        ("Columns up to 'Mortgage Summary' follow the reference research workbook; darker grey headers are application additions.", False),
        ("Every value's source, URL, retrieval time, raw value and quality is listed on the 'Sources' sheet.", False),
        ("", False),
        ("Colour legend", True),
        ("Entire row red (Montgomery): Tax Claim Bureau claim-year 2024 balance is $0.00 — 'Resolved / no 2024 balance' (unless overridden).", False),
        ("Orange valuation cell: value below the interest threshold; the property is marked Not interested.", False),
        ("Orange assessment cell (Delaware): assessment value below $100,000 (informational).", False),
        ("Green / red Tax Sale Status (Delaware): 'Listed for Upset Sale' / any other or missing status.", False),
        ("Mortgage Summary and Mortgage Info: red = no satisfaction found; green = satisfaction confirmed; yellow = potential match — review required / insufficient data.", False),
        ("Text styles: italic = estimate; blue = manually entered or overridden; brown italic = stale (daily status older than the refresh window).", False),
        ("", False),
        ("Sources that require a user action (disclaimer, login, human verification) are never automated; see 'User Action Tasks'.", False),
        ("Market values come only from licensed APIs or values the user recorded themselves; websites that prohibit scraping are not scraped.", False),
        ("", False),
        (DISCLAIMER, True),
    ]
    for i, (text, bold) in enumerate(lines, start=1):
        cell = legend.cell(i, 1, text)
        cell.font = Font(name="Arial", size=13 if i == 1 else 10, bold=bold)
        cell.alignment = Alignment(wrap_text=True)
    return wb, len(snaps)


def export_project(session: Session, project: Project, actor: str, job_id: int | None = None,
                   property_ids: list[int] | None = None) -> ExportRecord:
    wb, count = build_workbook(session, project, property_ids)
    settings = get_settings()
    folder = settings.exports_dir / str(project.id)
    folder.mkdir(parents=True, exist_ok=True)
    safe = re.sub(r"[^A-Za-z0-9._-]+", "_", project.name).strip("_")[:60] or "project"
    filename = f"{safe}_{utcnow():%Y%m%d_%H%M%S}.xlsx"
    path: Path = folder / filename
    wb.save(path)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    record = ExportRecord(project_id=project.id, job_id=job_id, filename=filename, file_path=str(path), sha256=digest,
                          row_count=count, created_by=actor)
    session.add(record)
    session.flush()
    record_audit(session, actor, "export.xlsx", "project", project.id, project.id,
                 {"export_id": record.id, "filename": filename, "rows": count, "sha256": digest})
    return record


__all__ = ["build_workbook", "export_project", "SourceLookup"]
