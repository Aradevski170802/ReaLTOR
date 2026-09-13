"""Inspect a reference workbook and capture its layout as a JSON template spec (no data values are kept).

Usage:  python -m app.export.template_inspector "<reference.xlsx>" app/export/templates/montco_reference.json
"""

from __future__ import annotations

import json
import sys
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter

TEMPLATES_DIR = Path(__file__).with_name("templates")


def _rgb(color) -> str | None:
    try:
        return color.rgb if color is not None and isinstance(color.rgb, str) else None
    except AttributeError:
        return None


def inspect_workbook(path: str | Path, sample_rows: int = 3) -> dict[str, Any]:
    wb = openpyxl.load_workbook(path)
    spec: dict[str, Any] = {"source_filename": Path(path).name, "sheets": []}
    for ws in wb.worksheets:
        header_row = 1
        columns = []
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(header_row, col)
            if cell.value is None:
                continue
            letter = get_column_letter(col)
            data_formats = sorted({ws.cell(r, col).number_format for r in range(2, min(ws.max_row, 1 + sample_rows) + 1)})
            columns.append({
                "letter": letter,
                "header": str(cell.value).strip(),
                "width": ws.column_dimensions[letter].width if letter in ws.column_dimensions else None,
                "header_fill": _rgb(cell.fill.fgColor) if cell.fill and cell.fill.fill_type else None,
                "header_font": {"name": cell.font.name, "size": cell.font.sz, "bold": cell.font.b, "color": _rgb(cell.font.color)},
                "header_alignment": {"horizontal": cell.alignment.horizontal, "vertical": cell.alignment.vertical, "wrap": cell.alignment.wrap_text},
                "data_number_formats": data_formats,
                "data_wrap": bool(ws.cell(2, col).alignment.wrap_text) if ws.max_row >= 2 else None,
                "data_font_size": ws.cell(2, col).font.sz if ws.max_row >= 2 else None,
            })
        fills: dict[str, int] = {}
        for row in ws.iter_rows(min_row=2, max_row=min(ws.max_row, 2000)):
            for c in row:
                if c.fill and c.fill.fill_type:
                    rgb = _rgb(c.fill.fgColor)
                    if rgb:
                        fills[rgb] = fills.get(rgb, 0) + 1
        spec["sheets"].append({
            "title": ws.title,
            "freeze_panes": ws.freeze_panes,
            "auto_filter": ws.auto_filter.ref,
            "header_row_height": ws.row_dimensions[1].height,
            "data_row_height": ws.row_dimensions[2].height if ws.max_row >= 2 else None,
            "row_count": max(0, ws.max_row - 1),
            "columns": columns,
            "fill_colors_used": dict(sorted(fills.items(), key=lambda kv: -kv[1])),
            "conditional_formatting": [
                {"range": str(rng.sqref), "type": rule.type, "formula": rule.formula}
                for rng, rules in ws.conditional_formatting._cf_rules.items() for rule in rules
            ],
            "text_rows": [ws.cell(r, 1).value for r in range(1, ws.max_row + 1)] if ws.max_column == 1 else None,
        })
    return spec


def load_builtin(county: str) -> dict[str, Any] | None:
    path = TEMPLATES_DIR / f"{county}_reference.json"
    return json.loads(path.read_text(encoding="utf-8")) if path.exists() else None


def header_style_lookup(spec: dict[str, Any] | None, sheet_title: str) -> dict[str, dict[str, Any]]:
    if not spec:
        return {}
    sheet = next((s for s in spec.get("sheets", []) if s["title"] == sheet_title), None)
    return {c["header"].strip().lower(): c for c in sheet["columns"]} if sheet else {}


def sheet_spec(spec: dict[str, Any] | None, sheet_title: str) -> dict[str, Any] | None:
    if not spec:
        return None
    return next((s for s in spec.get("sheets", []) if s["title"] == sheet_title), None)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit(__doc__)
    result = inspect_workbook(sys.argv[1])
    for sheet in result["sheets"]:
        sheet.pop("text_rows", None)  # layout only; no cell content is copied into the template
    out = Path(sys.argv[2])
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(result, indent=2), encoding="utf-8")
    print(f"Wrote {out}: " + ", ".join(f"{s['title']} ({len(s['columns'])} cols)" for s in result["sheets"]))
